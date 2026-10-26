"""Build the concise final LNCS manuscript from the supplied Word template and evidence.

Run with the bundled artifact Python; render with scripts/render_report.py on Windows.
The source document is never overwritten. SVG diagram sources accompany high-resolution PNGs.
"""

import html
import json
import math
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assignment_docs/updated"
FIGURES = OUT / "figures"
SOURCE = Path(r"C:\Users\aksha\Downloads\Cloud_Equipment_Monitoring_Springer_LNCS.docx")


class Figure:
    def __init__(self, name, height):
        self.name, self.width, self.height, self.scale = name, 560, height, 8
        self.image = Image.new("RGB", (self.width * self.scale, height * self.scale), "white")
        self.draw = ImageDraw.Draw(self.image)
        self.svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="560" height="{height}" viewBox="0 0 560 {height}"><rect width="560" height="100%" fill="white"/>']

    def text(self, x, y, text, size=12, bold=False):
        font_path = r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"
        font = ImageFont.truetype(font_path, round(size * self.scale))
        self.draw.text((x * self.scale, y * self.scale), text, font=font, fill="#202020", anchor="mm")
        self.svg.append(f'<text x="{x}" y="{y}" text-anchor="middle" dominant-baseline="central" font-family="Arial,sans-serif" font-size="{size}" font-weight="{"bold" if bold else "normal"}">{html.escape(text)}</text>')

    def box(self, x, y, w, h, lines):
        self.draw.rounded_rectangle((x * 8, y * 8, (x + w) * 8, (y + h) * 8), radius=24, fill="#f3f3f3", outline="#444444", width=8)
        self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="3" fill="#f3f3f3" stroke="#444"/>')
        for i, line in enumerate(lines):
            font = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf" if i == 0 else r"C:\Windows\Fonts\arial.ttf", 96)
            size = min(12, 12 * (w - 10) / max(1, font.getlength(line) / 8))
            self.text(x + w / 2, y + h / 2 + (i - (len(lines) - 1) / 2) * 15, line, size, i == 0)

    def arrow(self, points, label=None, at=None):
        scaled = [(x * 8, y * 8) for x, y in points]
        self.draw.line(scaled, fill="#444444", width=8)
        x, y = points[-1]
        dx, dy = x - points[-2][0], y - points[-2][1]
        angle = math.atan2(dy, dx)
        tips = [(x, y), (x - 7 * math.cos(angle - .45), y - 7 * math.sin(angle - .45)), (x - 7 * math.cos(angle + .45), y - 7 * math.sin(angle + .45))]
        self.draw.polygon([(a * 8, b * 8) for a, b in tips], fill="#444444")
        self.svg.append('<polyline fill="none" stroke="#444" points="' + ' '.join(f'{a},{b}' for a, b in points) + '"/>')
        self.svg.append('<polygon fill="#444" points="' + ' '.join(f'{a},{b}' for a, b in tips) + '"/>')
        if label and at:
            self.text(*at, label, 11)

    def save(self):
        FIGURES.mkdir(parents=True, exist_ok=True)
        self.image.save(FIGURES / (self.name + ".png"), optimize=True, dpi=(933, 933))
        (FIGURES / (self.name + ".svg")).write_text(''.join(self.svg) + '</svg>', encoding="utf-8")


def diagrams():
    f = Figure("architecture", 270)
    for x, w, lines in [(8, 112, ["EC2 gateway", "Seven assets"]), (155, 120, ["API Gateway", "Ingestion Lambda"]), (310, 90, ["SQS FIFO", "Telemetry"]), (435, 117, ["Processing", "Lambda"])]:
        f.box(x, 12, w, 48, lines)
    for start, end in [(120, 155), (275, 310), (400, 435)]:
        f.arrow([(start, 36), (end, 36)])
    f.box(145, 115, 120, 48, ["Inference Lambda", "Fault risk / action"])
    f.box(330, 115, 105, 48, ["DynamoDB", "Live / controls"])
    f.box(475, 115, 77, 48, ["S3", "History"])
    f.arrow([(470, 60), (470, 82), (205, 82), (205, 115)])
    f.arrow([(493, 60), (493, 95), (382, 95), (382, 115)])
    f.arrow([(520, 60), (520, 115)])
    f.arrow([(265, 139), (330, 139)], "Risk", (298, 126))
    f.box(330, 208, 105, 48, ["React dashboard", "Protected API"])
    f.arrow([(382, 163), (382, 208)])
    f.text(142, 206, "CloudFront / S3 serve React", 11)
    f.text(142, 225, "Cognito authenticates people", 11)
    f.text(510, 205, "Athena", 11, True)
    f.text(510, 222, "+ Glue", 11)
    f.arrow([(513, 163), (513, 188)])
    f.save()

    f = Figure("sequence", 324)
    centers = [43, 138, 232, 327, 422, 517]
    labels = [["EC2", "Gateway"], ["API / SQS", "Ingestion"], ["Processing", "Lambda"], ["DynamoDB", "State / controls"], ["Inference", "Lambda"], ["Dashboard", "via API"]]
    for x, lines in zip(centers, labels, strict=True):
        f.box(x - 41, 8, 82, 40, lines)
        f.draw.line([(x * 8, 48 * 8), (x * 8, 315 * 8)], fill="#aaaaaa", width=6)
        f.svg.append(f'<line x1="{x}" y1="48" x2="{x}" y2="315" stroke="#aaa"/>')
    for y, a, b, label in [(78, 0, 1, "Signed telemetry"), (112, 1, 2, "Consume batch"), (146, 2, 3, "Commit state and history"), (180, 2, 4, "Request inference"), (214, 4, 3, "Save risk / maintenance"), (248, 5, 3, "Read current status"), (282, 0, 3, "Poll controls; apply action")]:
        f.arrow([(centers[a], y), (centers[b], y)], label, ((centers[a] + centers[b]) / 2, y - 12))
    f.save()

    f = Figure("model_lifecycle", 172)
    items = [(8, ["Simulated runs", "Independent episodes"]), (150, ["Local training", "Validation / test"]), (292, ["Versioned S3", "Model + evaluation"]), (434, ["Inference Lambda", "Portable learned trees"])]
    for x, lines in items:
        f.box(x, 16, 118, 52, lines)
    for a, b in [(126, 150), (268, 292), (410, 434)]:
        f.arrow([(a, 42), (b, 42)])
    f.text(280, 109, "Telemetry windows → fault risk within five simulated minutes", 13, True)
    f.text(280, 139, "CloudFormation selects the version; a prior release supports rollback", 11)
    f.save()


def main():
    release = json.loads((ROOT / "models/release.json").read_text())
    evaluation = json.loads((ROOT / release["evaluationPath"]).read_text())
    live = json.loads((ROOT / "docs/predictive-live-evidence.json").read_text())
    assert len(live["scenarios"]) == 2, "Report requires completed live evidence"
    diagrams()
    d = Document(SOURCE)
    for child in list(d._element.body):
        if child.tag != qn("w:sectPr"):
            d._element.body.remove(child)
    section = d.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(5.2)
    section.left_margin = section.right_margin = Cm(4.4)
    for part in (section.header, section.footer):
        for p in part.paragraphs:
            p.text = ""
    for style in d.styles:
        if style.type == 1:
            style.font.name = "Times New Roman"
            style.font.color.rgb = RGBColor(0, 0, 0)
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(10)
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(.4)
    for name, size in [("Heading 1", 12), ("Heading 2", 10), ("Title", 14)]:
        st = d.styles[name]
        st.font.size, st.font.bold = Pt(size), True
        st.paragraph_format.first_line_indent = Cm(0)
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(12), Pt(6)
        st.paragraph_format.keep_with_next = True
        if st._element.pPr is not None:
            number = st._element.pPr.find(qn("w:numPr"))
            if number is not None:
                st._element.pPr.remove(number)
    d.core_properties.title = "Predictive Maintenance and Monitoring of Simulated Industrial Equipment on AWS"
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""
    d.core_properties.comments = ""
    d.core_properties.subject = "College cloud computing project"

    def paragraph(text, indent=True, size=10):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(.4 if indent else 0)
        p.paragraph_format.widow_control = True
        r = p.add_run(text)
        r.font.name, r.font.size = "Times New Roman", Pt(size)
        return p

    def heading(text, level=1):
        p = d.add_paragraph(text, style=f"Heading {level}")
        # Explicit numbering avoids inheriting template-specific numbering definitions.
        num = p._p.get_or_add_pPr().find(qn("w:numPr"))
        if num is not None:
            p._p.get_or_add_pPr().remove(num)
        return p

    def figure(name, caption, width=12.2):
        p = d.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(str(FIGURES / (name + ".png")), width=Cm(width))
        cap = paragraph(caption, False, 9)
        cap.paragraph_format.space_after = Pt(8)
        for inline in p._p.xpath('.//wp:docPr'):
            inline.set("descr", caption)

    def table(caption, headers, rows, widths):
        cap = paragraph(caption, False, 9)
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(4)
        t = d.add_table(rows=1, cols=len(headers))
        t.autofit = False
        for column, width in zip(t.columns, widths, strict=True):
            column.width = Cm(width)
        for cell, text, width in zip(t.rows[0].cells, headers, widths, strict=True):
            cell.text, cell.width = text, Cm(width)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell, text, width in zip(t.add_row().cells, row, widths, strict=True):
                cell.text, cell.width = str(text), Cm(width)
        for index, row in enumerate(t.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side in ("top", "bottom", "left", "right"):
                    margin = OxmlElement(f"w:{side}"); margin.set(qn("w:w"), "60"); margin.set(qn("w:type"), "dxa"); margins.append(margin)
                tcpr.append(margins)
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "bottom", "left", "right"):
                    b = OxmlElement(f"w:{side}"); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:color"), "D9D9D9"); borders.append(b)
                tcpr.append(borders)
                if index == 0:
                    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "EEEEEE"); tcpr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.keep_with_next = index == 0
                    for r in p.runs:
                        r.font.name, r.font.size, r.font.bold = "Times New Roman", Pt(9), index == 0
        d.add_paragraph().paragraph_format.space_after = Pt(2)

    title = d.add_paragraph(d.core_properties.title, style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(14)
    authors = ROOT / ".runtime/report-authors.json"
    if authors.exists():
        for line in json.loads(authors.read_text())["lines"]:
            paragraph(line, False).alignment = WD_ALIGN_PARAGRAPH.CENTER
    abstract = paragraph("Abstract. This paper presents a deployed AWS system for monitoring seven simulated industrial assets and predicting faults within the next five simulated minutes. One EC2 gateway represents three induction motors, a centrifugal pump, an air compressor, a cooling fan and a generator. Signed telemetry passes through API Gateway, Lambda and a FIFO queue into DynamoDB and an S3 archive. Configurable deterioration profiles produce thermal and mechanical faults. A dataset of 420 independent simulation runs supports local training and evaluation. The trained random forest is exported as portable learned trees and served by a separate inference Lambda. On 76 held-out runs, fault-warning precision was 90.3%, recall was 94.3%, and the median warning lead was five simulated minutes. The deployed system demonstrated warnings before faults, automatic cooldown and restart, and mechanical faults requiring operator repair. An authenticated React dashboard combines measurements, predictions, maintenance history and bounded historical reports. The results establish a reproducible teaching demonstration; they do not establish predictive accuracy on real industrial equipment.", False, 9)
    abstract.paragraph_format.left_indent = abstract.paragraph_format.right_indent = Cm(.5)
    abstract.paragraph_format.space_after = Pt(6)
    paragraph("Keywords: Predictive maintenance · Cloud computing · Industrial simulation · Random forest · Serverless inference · AWS", False, 9)

    heading("1 Introduction")
    paragraph("Equipment monitoring becomes more useful when it gives an operator time to act before a fault occurs. Condition-based maintenance links observed equipment condition to maintenance decisions [1]. A college demonstration must also make the data path, failure behavior and resulting actions visible without depending on physical sensors or an expensive deployment.", False)
    paragraph("We implemented a small industrial monitoring platform on AWS and extended it with predictive maintenance. The system combines a configurable simulator, retry-safe telemetry processing, learned fault warnings, two maintenance paths and a protected dashboard. Its objective is to demonstrate these functions together and support them with reproducible tests and measured evidence. The deployment uses the Sydney region and a USD 50 project budget.")
    paragraph("The contribution is an integrated teaching implementation rather than a new predictive algorithm. Scheduled deterioration makes demonstrations repeatable, while the classifier uses observed measurement windows rather than the simulator's hidden failure time. This distinction allows the project to show a learned warning before an ordinary threshold alert.")

    heading("2 Background and Design Rationale")
    paragraph("A threshold rule identifies a condition that has already crossed a limit. Prognostic methods instead use condition information to anticipate future deterioration [1]. Random forests combine multiple learned decision trees [2] and provide a compact classifier for the small tabular feature set used here. We use scikit-learn for training [3], then export the tree structures so that cloud inference does not require a scientific Python installation.", False)
    paragraph("Evaluation must respect dependencies between measurements. Randomly splitting overlapping windows would place closely related samples in training and testing and give an optimistic estimate. Scikit-learn's evaluation guidance explains the importance of temporal dependence and suitable splits [4]. We therefore separate entire simulation episodes before fitting the model or choosing its warning threshold. Independent episodes share the simulator's assumptions, so this split measures performance within that synthetic domain.")
    paragraph("The cloud design separates ingestion, processing and inference. AWS documents that Lambda consumers of SQS can receive messages more than once [5]. Event identifiers and conditional DynamoDB transactions [6] therefore protect telemetry state under retries, while prediction failure remains separate from the durable monitoring path.")

    heading("3 System Implementation")
    heading("3.1 Telemetry and Cloud Architecture", 2)
    paragraph("Figure 1 shows the deployed architecture. One EC2 instance runs the seven assets and sends a batch every ten seconds using its IAM role and signed API requests. API Gateway invokes an ingestion Lambda, which validates registered identities, timestamps, numeric bounds and schema fields. HTTP 202 means that the readings reached the FIFO queue, not that downstream work has completed.", False)
    figure("architecture", "Fig. 1. Main AWS data path. Inference adds fault risk and maintenance actions to live state; S3 history is queried through Athena and Glue")
    paragraph("The processing Lambda archives each batch as compressed, date-partitioned NDJSON in private S3. For each event, a DynamoDB transaction commits its deduplication marker, latest state, recent reading and any alert transition or notification intent. Older observations cannot replace newer state. A failed batch can be retried, and repeatedly failing messages move to a dead-letter queue. S3 and DynamoDB are not one transaction; stable archive identities and retry-safe writes handle partial completion.")
    paragraph("Figure 2 follows a reading through monitoring and prediction. Processing asynchronously invokes the predictor after durable telemetry work. A minute schedule also reconciles inference and maintenance. Both paths read a bounded recent window from current state rather than scanning the archive. Conditional writes prevent an older prediction from replacing a newer one. A separate scheduled outbox worker sends SNS alerts; SNS acceptance is distinguished from inbox receipt.")
    figure("sequence", "Fig. 2. Simplified telemetry, prediction and control sequence. Dashboard reads pass through the authenticated application API")

    heading("3.2 Equipment and Failure Profiles", 2)
    paragraph("The three original motors are induction motors. Four additional assets have distinct nominal temperatures, power requirements and supplementary measurements (Table 1). Common temperature, vibration and power fields preserve comparable charts and the existing reports. Generator output is a separate metric; its common power field represents auxiliary consumption.", False)
    table("Table 1. Simulated equipment and supplementary measurements", ["Equipment", "Count", "Supplementary readings"], [("Induction motor", 3, "Load and shaft speed"), ("Centrifugal pump", 1, "Flow and pressure"), ("Air compressor", 1, "Pressure"), ("Cooling fan", 1, "Airflow"), ("Generator", 1, "Output power, voltage, frequency")], [3.6, 1.2, 7.4])
    paragraph("Profiles select recoverable overheating, mechanical vibration, alternating modes or persistent overheating. A nominal 30-minute cycle can be configured between 10 and 120 simulated minutes. Seeded variation changes cycle timing by approximately ten percent; load and sensor noise vary within a run. Deterioration gradually affects measurements instead of switching directly from healthy to failed. Operators can disable deterioration or retain the manual overheat, vibration and disconnect controls.")
    paragraph("Presentation speeds of 5× and 10× advance simulated time faster while retaining ten-second transmissions. Consequently, five simulated minutes represent 30 wall-clock seconds at 10×. The dashboard displays the speed so that prediction lead time is not mistaken for elapsed real time. Mechanical damage latches the device into a stopped condition until an operator repairs or resets it.")

    heading("3.3 Predictive Model and Reproducible Releases", 2)
    paragraph("The training script generates 420 seeded episodes across equipment types, varied cycle lengths and demonstration speeds, including healthy episodes. It produced 57,032 pre-fault feature windows. Eight inputs describe current temperature and vibration, their trends per simulated minute, normalized power, load and recent measurement ranges. The window covers at most two simulated minutes. Absolute time since cycle start, profile settings and future fault times are excluded from the input vector.", False)
    paragraph("A sample is labeled overheating or vibration when the first corresponding threshold crossing occurs within the next five simulated minutes; otherwise it is labeled healthy. Post-fault samples are excluded from this warning evaluation. All windows from an episode remain together: 273 episodes train the model, 71 select the score threshold and 76 provide the held-out test. The classifier contains 32 trees, each limited to depth eight, with a minimum of 12 samples per leaf and balanced class weights.")
    paragraph("The warning score is one minus the forest's healthy-class score. Its validation-selected threshold is 0.80. The most likely fault class supplies the warning label. This is an uncalibrated classifier score, not a guaranteed failure probability. The exported inference implementation reproduces scikit-learn's float32 threshold comparisons; 500 held-out inputs passed a numerical parity check.")
    figure("model_lifecycle", "Fig. 3. Lightweight model lifecycle. Training and evaluation run locally; versioned artifacts support cloud deployment and rollback")
    paragraph("Figure 3 shows the release process. Training dependencies are pinned in an isolated local environment. A release must meet precision and recall gates of at least 0.75 before publication. The model and evaluation are stored under a versioned S3 prefix, and CloudFormation parameters select the deployed version. Promotion and rollback use the same deployment path. Inference runs in a 256 MiB Lambda using the exported trees and standard Python, avoiding a continuously running ML server.")

    heading("3.4 Maintenance and Dashboard", 2)
    paragraph("Three consecutive readings above 75 °C or 4.5 mm/s activate the corresponding fault; three healthy readings resolve it. Predictions provide earlier warnings while these deterministic rules remain the fault criteria. The interface uses Healthy, Warning and Fault labels, with offline connectivity shown separately. Equipment pages show type-specific readings, the five-minute outlook, model version and maintenance history. Operators can configure profiles or record repair; viewers can inspect data and reports.", False)
    paragraph("For recoverable overheating, the cloud controller records a cooldown request in the control table. The simulator stops the affected device, cools it toward its nominal temperature and restarts it after the bounded cooldown. A persistent thermal condition remains stopped and escalates to Repair needed; the controller does not repeatedly restart the same fault episode. Mechanical faults also require explicit operator action. Conditional control updates respect a superseding operator request. The history records predictions, cooldown, automatic restart and manual repair without adding a separate human verification workflow.")
    paragraph("Cognito provides hosted authorization-code login with PKCE. API Gateway protects application routes, and the backend checks operator membership before accepting controls or profile changes. Service roles supply cloud credentials. S3 buckets remain private, and CloudFront serves the React build. Recent readings are retained for 24 hours and archived telemetry for 30 days. Athena exposes three bounded reports for temperature, vibration breaches and running power, with Glue defining the archive schema. CloudWatch logs and alarms monitor platform health.")

    heading("4 Evaluation and Results")
    heading("4.1 Held-Out Prediction Results", 2)
    paragraph("Table 2 compares the forest with a causal baseline that extrapolates the current temperature and vibration trends to their fault thresholds. All reported warning metrics concern the next-five-minute binary decision on the held-out windows. The baseline uses the same window information. No test-set result was used to select the deployed model or score threshold.", False)
    table("Table 2. Warning performance on 76 held-out simulation episodes", ["Method", "Precision", "Recall", "F1"], [("Random forest", "0.903", "0.943", "0.923"), ("Linear trend baseline", "0.981", "0.814", "0.890")], [5.0, 2.4, 2.4, 2.4])
    paragraph("The forest improves warning recall and F1 at the cost of more false positives. Its false-positive rate on healthy test windows was 0.93%, compared with 0.14% for the baseline. Of the 76 test episodes, 51 reached a fault; each received at least one warning within the preceding five simulated minutes. The median lead was five minutes and the minimum was 1.67 minutes. The remaining episodes were healthy. These figures distinguish per-window errors from episode-level detection and do not imply that every individual warning was correct.")

    heading("4.2 Deployed Demonstrations and Verification", 2)
    paragraph("The live acceptance exercised a pump thermal episode and an air-compressor mechanical episode through the deployed API, EC2 simulator, queue, processing and inference functions. In both cases, a stored model warning preceded the active fault. The pump cooled and restarted automatically. The compressor stayed stopped with Repair needed until the authenticated operator issued a repair action. The maintenance history retained all five action categories, and a viewer's direct profile-change request was rejected. A separate live persistent-overheating test confirmed that the device remained stopped, escalated once and required manual repair. Model promotion and rollback were also verified through new live predictions. All seven assets reported online, including generator-specific output measurements.", False)
    table("Table 3. Verification evidence for the expanded implementation", ["Layer", "Result"], [("Backend and infrastructure", "82 local tests; schema and security rules passed"), ("React browser checks", "38 controlled scenarios and five real browser journeys"), ("Model export", "500 held-out predictions matched training-library output"), ("Live maintenance", "Predictive warning, automatic restart and manual repair")], [4.2, 8.0])
    paragraph("Five real browser journeys verified hosted PKCE login, live charts and mobile viewer access, automatic and manual maintenance controls, all three historical reports with CSV export, persistent team-role changes, and generator measurements with model details. Local failure-injection tests cover duplicate and older-event handling, invalid messages, storage retries and fault transitions. These checks support the defined application behavior; they are not a throughput benchmark or a guarantee of future cloud availability.")

    heading("4.3 Cost and Operating Scope", 2)
    paragraph("The implementation retains one small EC2 simulator and adds no second host, GPU, managed ML endpoint, NAT gateway or load balancer. Local training has no AWS training charge. Inference is event-driven, with a minute reconciliation schedule and bounded retries. Its cost depends on measured execution time and the existing database writes; the seven-device expansion increases telemetry storage and processing. Initial CloudWatch observations of 84 inference calls averaged 700 ms, with a maximum of 2,586 ms including cold and release transitions. A conservative 730-hour planning model estimates USD 1.08 for inference Lambda and USD 45.47 for the complete seven-device platform before credits. The estimate includes bounded database activity and variable-service allowances, and is not a measured bill. The project retains a USD 50 service-usage ceiling and notifications at USD 20, 30 and 40. Billing snapshots are delayed and account-wide, so they must not be presented as a precise project-only cost.", False)
    paragraph("No shutdown date or scheduled teardown is configured. The service remains online until the user requests shutdown. Bounded history, seven-day logs, on-demand database capacity and limited historical queries control ongoing usage.")

    heading("5 Limitations and Conclusion")
    paragraph("The principal limitation is synthetic validity. Training and test episodes come from the same simulator family, and the implemented fault dynamics are simplified teaching models. No real sensor dataset, independent physical failure process or industrial safety assessment was used. The model does not estimate remaining useful life or justify autonomous repair of real equipment. Restart-based recovery applies only to the explicitly modeled transient thermal scenario; mechanical repair resets simulated damage rather than repairing hardware.", False)
    paragraph("A single gateway and FIFO message group remain shared dependencies. Dashboard updates use polling, notification delivery can repeat after a worker crash, and recovery evidence covers a bounded archive exercise rather than regional disaster recovery. A useful next evaluation would introduce independently designed fault dynamics and real operating traces, measure drift and calibration, and benchmark latency and cost over a longer controlled workload.")
    paragraph("Within these limits, the project demonstrates predictive maintenance as an operational loop: configurable equipment deterioration produces telemetry, a locally trained and versioned model warns before a fault, and the cloud coordinates automatic or manual maintenance. The deployed dashboard makes the measurements, outlook and actions understandable while preserving the existing authentication, history and reporting functions. The result is a reproducible college-scale cloud and ML operations demonstration supported by held-out evaluation and live evidence.")

    heading("References")
    references = [
        "Jardine, A.K.S., Lin, D., Banjevic, D.: A review on machinery diagnostics and prognostics implementing condition-based maintenance. Mechanical Systems and Signal Processing 20(7), 1483–1510 (2006). https://doi.org/10.1016/j.ymssp.2005.09.012",
        "Breiman, L.: Random forests. Machine Learning 45, 5–32 (2001). https://doi.org/10.1023/A:1010933404324",
        "Scikit-learn Developers: RandomForestClassifier documentation. https://scikit-learn.org/stable/modules/generated/sklearn.ensemble.RandomForestClassifier.html. Accessed 27 Oct 2026",
        "Scikit-learn Developers: Cross-validation: evaluating estimator performance. https://scikit-learn.org/stable/modules/cross_validation.html. Accessed 27 Oct 2026",
        "Amazon Web Services: Using Lambda with Amazon SQS. https://docs.aws.amazon.com/lambda/latest/dg/with-sqs.html. Accessed 27 Oct 2026",
        "Amazon Web Services: Amazon DynamoDB transactions: how it works. https://docs.aws.amazon.com/amazondynamodb/latest/developerguide/transaction-apis.html. Accessed 27 Oct 2026",
    ]
    for index, text in enumerate(references, 1):
        p = paragraph(f"{index}. {text}", False, 9)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(4)
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / "Cloud_Equipment_Monitoring_LNCS.docx"
    d.save(path)
    print(json.dumps({"docx": str(path), "words": sum(len(p.text.split()) for p in d.paragraphs), "figures": len(d.inline_shapes), "tables": len(d.tables), "modelVersion": evaluation["version"], "authorDetailsProvided": authors.exists()}, indent=2))


if __name__ == "__main__":
    main()
